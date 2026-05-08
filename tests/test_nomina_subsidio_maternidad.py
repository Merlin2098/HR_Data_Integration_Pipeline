from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
import polars as pl

from src.modules.nomina.steps.step1_consolidar_planillas import (
    SUBSIDIO_MATERNIDAD_COL,
    asegurar_columna_subsidio_maternidad,
    leer_archivo_planilla,
)


def test_asegura_subsidio_maternidad_when_header_is_missing():
    df = pl.DataFrame({"DNI/CEX": ["1", "2"]})

    result = asegurar_columna_subsidio_maternidad(df)

    assert SUBSIDIO_MATERNIDAD_COL in result.columns
    assert result[SUBSIDIO_MATERNIDAD_COL].to_list() == [None, None]


def test_asegura_subsidio_maternidad_matches_header_case_insensitive():
    df = pl.DataFrame({"subsidio maternidad": [10.5, None]})

    result = asegurar_columna_subsidio_maternidad(df)

    assert "subsidio maternidad" not in result.columns
    assert result[SUBSIDIO_MATERNIDAD_COL].to_list() == [10.5, None]


def test_leer_archivo_planilla_descarta_filas_de_totales_y_subtotales(tmp_path: Path):
    workbook_path = tmp_path / "METSO_Planilla 2026-04 Empleados.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Planilla"

    ws["A3"] = "Reporte de Nomina por Periodo"
    ws["A4"] = "Ano: 2026 - Periodo: 4"

    headers = [
        "N°",
        "CÓDIGO",
        "CÓDIGO CORPORATIVO",
        "APELLIDO PATERNO",
        "APELLIDO MATERNO",
        "NOMBRES",
        "TIPO DE DOCUMENTO",
        "DNI/CEX",
        "NETO",
    ]
    for idx, header in enumerate(headers, start=1):
        ws.cell(row=6, column=idx, value=header)

    empleado = [
        1,
        "77071019",
        "445341",
        "ABANTO",
        "GALVAN",
        "ARIANA NICOLE",
        "DNI",
        "77071019",
        1250.5,
    ]
    total = ["+ INGRESOS DEL MES", None, None, 24, None, None, None, None, 841894.06]

    for idx, value in enumerate(empleado, start=1):
        ws.cell(row=7, column=idx, value=value)

    for idx, value in enumerate(total, start=1):
        ws.cell(row=8, column=idx, value=value)

    wb.save(workbook_path)

    result = leer_archivo_planilla(workbook_path, "2026-04")

    assert result.height == 1
    assert result["N°"].to_list() == [1]
    assert result["DNI/CEX"].to_list() == ["77071019"]
